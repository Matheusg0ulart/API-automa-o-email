import csv
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class EmailReporter:
    """Gera relatórios consolidados em Excel (.xlsx) e CSV com formatação profissional."""

    def __init__(self, reports_dir: Path):
        self.reports_dir = reports_dir
        self.reports_dir.mkdir(exist_ok=True)

    def generate_report(self, processed_records: List[Dict[str, Any]], dry_run: bool = False) -> Dict[str, Any]:
        """
        Salva o histórico de processamento em arquivos .xlsx e .csv.
        Retorna os caminhos dos arquivos e estatísticas resumidas.
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        prefix = "simulacao_" if dry_run else "processamento_"
        xlsx_path = self.reports_dir / f"{prefix}emails_{timestamp}.xlsx"
        csv_path = self.reports_dir / f"{prefix}emails_{timestamp}.csv"

        # 1. Estatísticas por Categoria
        stats: Dict[str, int] = {}
        for item in processed_records:
            cat = item.get("category_name", "Desconhecido")
            stats[cat] = stats.get(cat, 0) + 1

        # 2. Exportação CSV com UTF-8-SIG para compatibilidade com Excel
        headers = [
            "UID",
            "Data",
            "De (Nome)",
            "De (E-mail)",
            "Assunto",
            "Categoria",
            "Pasta Destino",
            "Score",
            "Urgência",
            "Valor Estimado",
            "Resumo IA",
            "Motivos Detectados",
            "Auto-Resposta",
            "Status"
        ]

        with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(headers)
            for item in processed_records:
                reasons_str = ", ".join(item.get("reasons", []))
                writer.writerow([
                    item.get("uid", ""),
                    item.get("date", ""),
                    item.get("sender_name", ""),
                    item.get("sender_email", ""),
                    item.get("subject", ""),
                    item.get("category_name", ""),
                    item.get("target_folder", ""),
                    item.get("score", 0),
                    item.get("urgency", "Média"),
                    item.get("estimated_value", "Não informado"),
                    item.get("summary", ""),
                    reasons_str,
                    "Enviada" if item.get("auto_replied") else "Não",
                    item.get("status", "")
                ])

        # 3. Exportação Excel com Estilização
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "E-mails Processados"

        # Cores e Estilos
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3")
        )

        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx, item in enumerate(processed_records, start=2):
            reasons_str = ", ".join(item.get("reasons", []))
            row_data = [
                item.get("uid", ""),
                item.get("date", ""),
                item.get("sender_name", ""),
                item.get("sender_email", ""),
                item.get("subject", ""),
                item.get("category_name", ""),
                item.get("target_folder", ""),
                item.get("score", 0),
                item.get("urgency", "Média"),
                item.get("estimated_value", "Não informado"),
                item.get("summary", ""),
                reasons_str,
                "Enviada" if item.get("auto_replied") else "Não",
                item.get("status", "")
            ]
            ws.append(row_data)

            # Estilização das linhas de dados
            fill_color = "F9FBFD" if row_idx % 2 == 0 else "FFFFFF"
            row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = row_fill
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)
                if col_idx in [1, 8, 9, 13, 14]:  # UID, Score, Urgência, Auto-Resposta, Status centralizados
                    cell.alignment = Alignment(horizontal="center")

        # Ajusta automaticamente a largura das colunas
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)

        wb.save(xlsx_path)

        return {
            "xlsx_path": str(xlsx_path),
            "csv_path": str(csv_path),
            "total_processed": len(processed_records),
            "categories_breakdown": stats
        }
